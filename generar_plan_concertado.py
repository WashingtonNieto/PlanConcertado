#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Actas de PLAN CONCERTADO  -  SENA / Articulacion con la Media
Formato GOR-F-084 Acta V02

Toma las competencias, RAP, guias, temas y actividades desde el archivo
RESULTADOS_A_EVALUAR_10_11_GUIA_TEMA_ACTIVIDAD.xlsx y los escribe sobre la
plantilla oficial en Word, conservando el formato del documento.

Uso:
    python generar_plan_concertado.py          -> abre la ventana

Requisitos:
    pip install python-docx openpyxl

Autor: generado para el instructor Washington Nieto
"""

import copy
import json
import os
import re
import sys
import unicodedata
from datetime import date, datetime

try:
    import openpyxl
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Twips
    from docx.table import _Row
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    print("Faltan librerias. Ejecute:  pip install python-docx openpyxl")
    raise

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "config_colegios.json")

ORDINALES = {1: "PRIMER", 2: "SEGUNDO", 3: "TERCER", 4: "CUARTO"}
NUM_PALABRA = {1: "una", 2: "dos", 3: "tres", 4: "cuatro", 5: "cinco", 6: "seis",
               7: "siete", 8: "ocho", 9: "nueve", 10: "diez", 11: "once",
               12: "doce", 13: "trece", 14: "catorce", 15: "quince"}
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# Posiciones dentro de la tabla de la plantilla (tabla 0)
R_ACTA = 0
R_NOMBRE = 1
R_CIUDAD = 2
R_LUGAR = 3
R_AGENDA = 4
R_OBJETIVO = 5
R_DESARROLLO = 7
R_CONCLUSIONES = 9
R_COMPROMISO_INI = 12
R_COMPROMISO_FIN = 16
R_INSTRUCTOR = 19
R_VOCERO = 20
R_APRENDICES = 21

# Sangrias de las listas, en twips (1 cm = 567).
# nivel: (sangria izquierda, sangria francesa)
#
# Cada formato es (tipo, simbolo, fuente del simbolo).
#
# DESARROLLO DE LA REUNIÓN:   1.  /  a.  /  i.
IND_NIVEL = {0: (312, 236), 1: (672, 240), 2: (1032, 240)}
FORMATOS_NIVEL = [("decimal", "%1.", None), ("lowerLetter", "%2.", None),
                  ("lowerRoman", "%3.", None)]

# AGENDA O PUNTOS PARA DESARROLLAR:   1.  /  •
IND_AGENDA = {0: (312, 236), 1: (1134, 284)}
FORMATOS_AGENDA = [("decimal", "%1.", None), ("bullet", "", "Symbol")]

# CONCLUSIONES:   -  /  o
IND_CONCLUSIONES = {0: (312, 236), 1: (1134, 284)}
FORMATOS_CONCLUSIONES = [("bullet", "-", "Calibri"), ("bullet", "o", "Courier New")]


# --------------------------------------------------------------------------
# Configuracion
# --------------------------------------------------------------------------
def cargar_config(ruta=CONFIG_PATH):
    with open(ruta, encoding="utf-8") as fh:
        return json.load(fh)


# --------------------------------------------------------------------------
# Lectura del Excel
# --------------------------------------------------------------------------
def _norm(texto):
    """Quita tildes y pasa a mayusculas, para comparar nombres de hoja."""
    if texto is None:
        return ""
    sin = unicodedata.normalize("NFKD", str(texto))
    sin = "".join(c for c in sin if not unicodedata.combining(c))
    return sin.upper().strip()


def _hoja_de_grado(wb, grado):
    """Devuelve la hoja 'RESULTADOS A EVALUAR 10' u 11, tolerando variaciones."""
    grado = str(grado).replace("°", "").strip()
    for ws in wb.worksheets:
        nombre = _norm(ws.title).replace("°", "")
        if "RESULTADOS" in nombre and re.search(r"\b%s\b" % grado, nombre):
            return ws
    raise ValueError("No se encontro la hoja de resultados para el grado %s." % grado)


def _fechas_trimestre(texto):
    """De 'TRIMESTRE 2 Inicio: 20/04/2026 Fin: 06/07/2026' saca las dos fechas."""
    ini = re.search(r"Inicio\s*:\s*(\d{1,2}/\d{1,2}/\d{4})", texto or "", re.I)
    fin = re.search(r"Fin\s*:\s*(\d{1,2}/\d{1,2}/\d{4})", texto or "", re.I)
    return (ini.group(1) if ini else ""), (fin.group(1) if fin else "")


def _fecha_larga(cadena):
    """'20/04/2026' -> '20 de abril de 2026'. Si no puede, devuelve tal cual."""
    try:
        d = datetime.strptime(cadena.strip(), "%d/%m/%Y").date()
        return "%d de %s de %d" % (d.day, MESES[d.month - 1], d.year)
    except Exception:
        return cadena


def _limpiar(texto):
    if texto is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(texto)).strip()


def a_oracion(texto, terminos=None):
    """
    Pasa un texto escrito TODO EN MAYUSCULAS a formato de oracion:
    minusculas con la primera letra de cada oracion en mayuscula, respetando
    los terminos tecnicos (SQL, MySQL, HTML5...) del archivo de configuracion.

    Si el texto ya viene redactado normalmente, lo devuelve sin tocar.
    """
    texto = _limpiar(texto)
    letras = [c for c in texto if c.isalpha()]
    if not letras:
        return texto
    proporcion = sum(1 for c in letras if c.isupper()) / float(len(letras))
    if proporcion < 0.6:                     # ya esta en formato normal
        return texto

    resultado = texto.lower()
    # Mayuscula inicial de cada oracion. En espanol despues de dos puntos va
    # minuscula, por eso solo se considera el punto y el salto de linea.
    # El primer grupo salta lo que anteceda a la primera letra ("1) ", "- "...).
    resultado = re.sub(
        r"(^[^a-zA-ZáéíóúüñÁÉÍÓÚÜÑ]*|\.\s+|\n\s*)([a-záéíóúüñ])",
        lambda m: m.group(1) + m.group(2).upper(), resultado)
    # Restaurar la grafia correcta de los terminos tecnicos.
    for clave in sorted(terminos or {}, key=len, reverse=True):
        resultado = re.sub(r"\b%s\b" % re.escape(clave), terminos[clave],
                           resultado, flags=re.IGNORECASE)
    return resultado


def parse_actividades(texto):
    """Separa '1) uno. 2) dos.' en una lista y renumera desde 1."""
    if not texto:
        return []
    partes = re.split(r"\d+\)\s*", str(texto))
    items = []
    for p in partes:
        p = _limpiar(p.replace("\n", " ")).strip(" .;-")
        if p:
            items.append(p[0].upper() + p[1:] + ".")
    return items


def leer_trimestre(ruta_excel, grado, trimestre):
    """
    Devuelve un dict con la informacion del trimestre pedido:
    {inicio, fin, competencias:[{comp, raps:[...]}], guias:[...], temas:[...], actividades:[...]}
    """
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws = _hoja_de_grado(wb, grado)

    filas, actual = [], None
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not any(fila):
            continue
        trim_txt, fase, comp, rap, guia, tema, act = (list(fila) + [None] * 7)[:7]
        if trim_txt:
            actual = _limpiar(trim_txt.replace("\n", " "))
        filas.append({
            "trim": actual, "fase": _limpiar(fase), "comp": _limpiar(comp),
            "rap": rap or "", "guia": _limpiar(guia),
            "tema": _limpiar(tema), "act": act or "",
        })

    buscado = "TRIMESTRE %s" % trimestre
    propias = [f for f in filas if f["trim"] and f["trim"].upper().startswith(buscado)]
    if not propias:
        raise ValueError("El Excel no tiene datos para el %s en grado %s." % (buscado, grado))

    inicio, fin = _fechas_trimestre(propias[0]["trim"])

    competencias, guias, temas, actividades = [], [], [], []
    for f in propias:
        raps = [_limpiar(r) for r in str(f["rap"]).split("\n") if _limpiar(r)]
        existente = next((c for c in competencias if c["comp"] == f["comp"]), None)
        if existente:
            for r in raps:
                if r not in existente["raps"]:
                    existente["raps"].append(r)
        else:
            competencias.append({"comp": f["comp"], "raps": raps})
        if f["guia"] and f["guia"] not in guias:
            guias.append(f["guia"])
        if f["tema"] and f["tema"] not in temas:
            temas.append(f["tema"])
        actividades.extend(parse_actividades(f["act"]))

    return {
        "inicio": inicio, "fin": fin,
        "competencias": competencias,
        "guias": guias, "temas": temas, "actividades": actividades,
    }


# --------------------------------------------------------------------------
# Utilidades de Word (conservando el formato de la plantilla)
# --------------------------------------------------------------------------
def set_parrafo(par, texto, negrita=None):
    """Cambia el texto de un parrafo conservando la fuente del primer run."""
    runs = par.runs
    if not runs:
        run = par.add_run(texto)
    else:
        runs[0].text = texto
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
        run = runs[0]
    if negrita is not None:
        run.bold = negrita
    return par


class Lista(object):
    """Una numeracion creada en numbering.xml, con sus sangrias por nivel."""

    def __init__(self, num_id, sangrias):
        self.num_id = num_id
        self.sangrias = sangrias


def crear_lista(doc, formatos, sangrias):
    """
    Agrega a numbering.xml una numeracion nueva con los formatos indicados
    (decimal, lowerLetter, lowerRoman, bullet...) y devuelve un objeto Lista.
    Al ser nueva empieza en 1 y no altera las listas que ya trae la plantilla.
    """
    numeracion = doc.part.numbering_part.element
    abstractos = [int(a.get(qn("w:abstractNumId")))
                  for a in numeracion.findall(qn("w:abstractNum"))]
    numeros = [int(n.get(qn("w:numId"))) for n in numeracion.findall(qn("w:num"))]
    id_abstracto = max(abstractos or [0]) + 1
    id_num = max(numeros or [0]) + 1

    niveles = ""
    for i, (formato, texto, fuente) in enumerate(formatos):
        izq, colgante = sangrias[i]
        # Cada viñeta se dibuja con su fuente: Symbol para •, Courier New para o.
        rpr = ""
        if fuente:
            rpr = ('<w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s" w:cs="%s" '
                   'w:hint="default"/></w:rPr>' % (fuente, fuente, fuente))
        niveles += (
            '<w:lvl w:ilvl="%d"><w:start w:val="1"/><w:numFmt w:val="%s"/>'
            '<w:lvlText w:val="%s"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="%d" w:hanging="%d"/></w:pPr>%s</w:lvl>'
            % (i, formato, texto, izq, colgante, rpr))
    # Niveles restantes hasta 8: Word exige que existan aunque no se usen.
    for i in range(len(formatos), 9):
        niveles += (
            '<w:lvl w:ilvl="%d"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%%%d."/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="%d" w:hanging="240"/></w:pPr></w:lvl>'
            % (i, i + 1, 1440 + 360 * (i - 2)))

    abstracto = parse_xml(
        '<w:abstractNum %s w:abstractNumId="%d">'
        '<w:multiLevelType w:val="hybridMultilevel"/>%s</w:abstractNum>'
        % (nsdecls("w"), id_abstracto, niveles))
    numero = parse_xml('<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>'
                       % (nsdecls("w"), id_num, id_abstracto))

    existentes = numeracion.findall(qn("w:num"))
    if existentes:
        existentes[0].addprevious(abstracto)   # los abstractNum van antes
    else:
        numeracion.append(abstracto)
    numeracion.append(numero)
    return Lista(id_num, sangrias)


def aplicar_nivel(par, lista, nivel):
    """Pone (o quita) la vinieta y la sangria del parrafo."""
    ppr = par._p.get_or_add_pPr()
    if nivel is None:
        ppr._remove_numPr()
        par.paragraph_format.left_indent = Twips(lista.sangrias[0][0])
        par.paragraph_format.first_line_indent = Twips(0)
        return
    numpr = ppr.get_or_add_numPr()
    numpr.get_or_add_ilvl().val = nivel
    numpr.get_or_add_numId().val = lista.num_id
    izq, colgante = lista.sangrias[nivel]
    par.paragraph_format.left_indent = Twips(izq)
    par.paragraph_format.first_line_indent = Twips(-colgante)


def set_lineas(celda, lineas, desde=0, negritas=(), niveles=None, lista=None,
               izquierda=()):
    """
    Reescribe la celda desde el parrafo 'desde' con la lista de lineas dada,
    clonando ese parrafo como molde para conservar fuente, tamanio y sangria.

    'niveles' (opcional) es una lista del mismo largo que 'lineas' con el nivel
    de vinieta de cada una; None en una posicion = parrafo sin vinieta.
    'izquierda' son los indices que se alinean a la izquierda en vez de
    justificados (util para textos largos en mayusculas, que al justificarse
    quedan con huecos enormes entre palabras).
    """
    parrafos = celda.paragraphs
    if desde >= len(parrafos):
        desde = len(parrafos) - 1
    molde = copy.deepcopy(parrafos[desde]._element)
    for p in parrafos[desde:]:
        p._element.getparent().remove(p._element)
    for i, linea in enumerate(lineas):
        elemento = copy.deepcopy(molde)
        celda._tc.append(elemento)
        par = Paragraph(elemento, celda)
        set_parrafo(par, linea, negrita=True if i in negritas else None)
        if niveles is not None and lista is not None:
            aplicar_nivel(par, lista, niveles[i])
        if i in izquierda:
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT


def celda(tabla, fila, col):
    return tabla.rows[fila].cells[col]


def escribir_compromisos(tabla, actividades, ficha):
    """Ajusta el bloque de compromisos al numero real de actividades."""
    # Se capturan los elementos XML ANTES de borrar: los indices se corren.
    viejas = [tabla.rows[i]._tr
              for i in range(R_COMPROMISO_INI, R_COMPROMISO_FIN + 1)]
    molde = copy.deepcopy(viejas[0])
    ancla = tabla.rows[R_COMPROMISO_INI - 1]._tr     # fila de encabezados
    cuerpo = ancla.getparent()

    for tr in viejas:
        cuerpo.remove(tr)

    for texto in actividades:
        tr = copy.deepcopy(molde)
        ancla.addnext(tr)
        ancla = tr
        celdas = _Row(tr, tabla).cells
        # La actividad y el responsable van alineados a la izquierda: son
        # textos cortos en mayusculas que justificados quedan con huecos.
        par = set_parrafo(celdas[0].paragraphs[0], texto)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_parrafo(celdas[2].paragraphs[0], "")
        par = set_parrafo(celdas[5].paragraphs[0], "Aprendices Ficha N° %s" % ficha)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_parrafo(celdas[8].paragraphs[0], "Se anexa lista de los aprendices")


# --------------------------------------------------------------------------
# Generacion del acta
# --------------------------------------------------------------------------
def nombre_archivo(colegio, grado, ficha, trimestre):
    limpio = re.sub(r"[^\w\s-]", "", _norm(colegio)).strip()
    limpio = re.sub(r"\s+", " ", limpio)
    return "PLAN CONCERTADO T%s - %s - FICHA %s - GRADO %s.docx" % (
        trimestre, limpio, ficha or "SIN FICHA", grado)


def generar_acta(cfg, datos):
    """
    datos: dict con colegio, ambiente, grado, ficha, trimestre, vocero,
           instructor, fecha (date), hora_inicio, hora_fin, dia, acta_no
    Devuelve la ruta del archivo generado.
    """
    ruta_plantilla = os.path.join(BASE_DIR, cfg["archivo_plantilla"])
    ruta_excel = os.path.join(BASE_DIR, cfg["archivo_excel"])
    salida_dir = os.path.join(BASE_DIR, cfg.get("carpeta_salida", "Actas generadas"))
    os.makedirs(salida_dir, exist_ok=True)

    trimestre = int(datos["trimestre"])
    grado = str(datos["grado"]).replace("°", "").strip()
    ficha = _limpiar(datos.get("ficha"))
    ord_txt = ORDINALES.get(trimestre, str(trimestre))

    info = leer_trimestre(ruta_excel, grado, trimestre)
    terminos = cfg.get("terminos_tecnicos", {})
    doc = Document(ruta_plantilla)
    tabla = doc.tables[0]

    # --- Encabezado --------------------------------------------------------
    acta_no = _limpiar(datos.get("acta_no"))
    set_parrafo(celda(tabla, R_ACTA, 0).paragraphs[0],
                "ACTA No. %s" % (acta_no if acta_no else "…"))

    set_parrafo(
        celda(tabla, R_NOMBRE, 0).paragraphs[1],
        "PLAN CONCERTADO %s TRIMESTRE PROGRAMA: %s - FICHA N° %s Grado %s°." % (
            ord_txt, cfg["programa"], ficha or "………", grado))

    fecha = datos.get("fecha") or date.today()
    set_parrafo(celda(tabla, R_CIUDAD, 3).paragraphs[0],
                "%s, %d de %s de %d" % (cfg["ciudad"], fecha.day,
                                        MESES[fecha.month - 1], fecha.year))
    set_lineas(celda(tabla, R_CIUDAD, 7), [_limpiar(datos.get("hora_inicio"))], desde=1)
    set_lineas(celda(tabla, R_CIUDAD, 9), [_limpiar(datos.get("hora_fin"))], desde=1)

    set_lineas(celda(tabla, R_LUGAR, 0), [_limpiar(datos["colegio"])], desde=1)
    set_parrafo(celda(tabla, R_LUGAR, 3).paragraphs[0],
                "AMBIENTE DE FORMACIÓN %s" % _limpiar(datos.get("ambiente")))
    set_lineas(celda(tabla, R_LUGAR, 7),
               ["DIRECCIÓN / REGIONAL / CENTRO: %s" % cfg["centro"]], desde=0)

    # --- Agenda:  1.  con las guias como viñetas redondas en el punto 3 -----
    agenda = [
        ("Presentación del cronograma de formación para el %s trimestre "
         "(Inicio: %s – Fin: %s)." % (ord_txt.lower(), info["inicio"], info["fin"]), 0),
        ("Socialización de las competencias y Resultados de Aprendizaje (RAP) "
         "a desarrollar.", 0),
        ("Desglose de las Guías de Aprendizaje:", 0),
    ]
    primera_guia = len(agenda)
    agenda.extend((guia, 1) for guia in info["guias"])
    guias_agenda = set(range(primera_guia, len(agenda)))
    agenda.extend([
        ("Definición de actividades, entregables y metodología de evaluación.", 0),
        ("Establecimiento de compromisos y fechas límite.", 0),
    ])
    lista_agenda = crear_lista(doc, FORMATOS_AGENDA, IND_AGENDA)
    set_lineas(celda(tabla, R_AGENDA, 0), [a[0] for a in agenda], desde=1,
               niveles=[a[1] for a in agenda], lista=lista_agenda,
               izquierda=guias_agenda)

    # --- Objetivo ----------------------------------------------------------
    set_lineas(celda(tabla, R_OBJETIVO, 0), [
        "Establecer y concertar el plan de trabajo para el %s trimestre del programa "
        "Técnico en Programación de Software, garantizando que los aprendices conozcan "
        "las competencias, los resultados de aprendizaje esperados, la metodología de "
        "formación y los criterios de evaluación para el cumplimiento exitoso del "
        "programa en el %s." % (ord_txt.lower(), _limpiar(datos["colegio"])),
    ], desde=1)

    # --- Desarrollo de la reunion -----------------------------------------
    # Estructura de viñetas:  1. parrafo guia   a. competencia   i. RAP
    n_comp = len(info["competencias"])
    if n_comp == 1:
        encabezado = ("Durante la sesión se presentó la competencia principal para "
                      "este trimestre, con sus respectivos resultados de aprendizaje:")
    else:
        encabezado = ("Durante la sesión se presentaron las %s competencias principales "
                      "para este trimestre, con sus respectivos resultados de "
                      "aprendizaje:" % NUM_PALABRA.get(n_comp, str(n_comp)))

    bloque = [(encabezado, 0, False), ("", None, False)]
    for c in info["competencias"]:
        bloque.append((c["comp"], 1, True))                 # a.  en negrita
        bloque.extend((rap, 2, False) for rap in c["raps"])  # i.  ii.  iii.
        bloque.append(("", None, False))

    bloque.append(("Se enfatizó que el proceso de formación se divide en los "
                   "siguientes ejes fundamentales:", 0, False))
    # Los temas vienen del Excel en mayusculas sostenidas: se pasan a formato
    # de oracion y se alinean a la izquierda (justificados quedan con huecos).
    primer_tema = len(bloque)
    bloque.extend((a_oracion(tema, terminos), 1, False) for tema in info["temas"])
    ejes = set(range(primer_tema, len(bloque)))
    bloque.append(("", None, False))

    bloque.append(("Se acordó una metodología de aprendizaje basada en el análisis del "
                   "desempeño, donde las evidencias se recolectarán a través de talleres "
                   "prácticos, listas de chequeo, evidencias de producto y cuestionarios "
                   "de conocimiento.", 0, False))

    lista_desarrollo = crear_lista(doc, FORMATOS_NIVEL, IND_NIVEL)
    set_lineas(celda(tabla, R_DESARROLLO, 0),
               [b[0] for b in bloque], desde=0,
               negritas={i for i, b in enumerate(bloque) if b[2]},
               niveles=[b[1] for b in bloque], lista=lista_desarrollo,
               izquierda=ejes)

    # --- Conclusiones:  -  con las guias como sub-viñetas  o  --------------
    n_act = len(info["actividades"])
    conclusiones = [
        ("Se ratifica el inicio de la formación el %s y su finalización el %s, "
         "con sesiones los %s de %s a %s." % (
             _fecha_larga(info["inicio"]), _fecha_larga(info["fin"]),
             _limpiar(datos.get("dia")) or "las sesiones programadas",
             _limpiar(datos.get("hora_inicio")), _limpiar(datos.get("hora_fin"))), 0),
        ("Los aprendices aceptan el plan de %s actividades principales "
         "correspondientes a:" % NUM_PALABRA.get(n_act, str(n_act)), 0),
    ]
    primera_guia = len(conclusiones)
    conclusiones.extend((guia, 1) for guia in info["guias"])
    guias_concl = set(range(primera_guia, len(conclusiones)))
    conclusiones.append(
        ("Todo entregable debe ser remitido en formato PDF y nombrado según las "
         "convenciones establecidas (ej. ACTIVIDAD DE APRENDIZAJE N°….. – "
         "GUIA N°……).", 0))
    lista_concl = crear_lista(doc, FORMATOS_CONCLUSIONES, IND_CONCLUSIONES)
    set_lineas(celda(tabla, R_CONCLUSIONES, 0), [c[0] for c in conclusiones],
               desde=0, niveles=[c[1] for c in conclusiones], lista=lista_concl,
               izquierda=guias_concl)

    # --- Asistentes (antes de tocar los compromisos, que corren las filas) --
    programa_titulo = cfg.get("programa_titulo", cfg["programa"])
    set_parrafo(celda(tabla, R_INSTRUCTOR, 0).paragraphs[0],
                _limpiar(datos.get("instructor")))
    set_parrafo(celda(tabla, R_VOCERO, 0).paragraphs[0], _limpiar(datos.get("vocero")))
    par = set_parrafo(celda(tabla, R_APRENDICES, 1).paragraphs[0],
                      "Integrantes de la Ficha N° %s del programa %s"
                      % (ficha or "………", programa_titulo))
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # --- Compromisos: actividades renumeradas y en formato de oracion ------
    actividades = ["%d) %s" % (i + 1, a_oracion(t, terminos))
                   for i, t in enumerate(info["actividades"])]
    escribir_compromisos(tabla, actividades, ficha or "………")

    ruta = os.path.join(salida_dir, nombre_archivo(
        datos["colegio"], grado, ficha, trimestre))
    doc.save(ruta)
    return ruta


# --------------------------------------------------------------------------
# Interfaz grafica (Tkinter)
# --------------------------------------------------------------------------
def main():
    import tkinter as tk
    from tkinter import messagebox, ttk

    cfg = cargar_config()

    ventana = tk.Tk()
    ventana.title("Plan Concertado - SENA | Articulación con la Media")
    ventana.resizable(False, False)
    try:
        ventana.call("tk", "scaling", 1.2)
    except Exception:
        pass

    marco = ttk.Frame(ventana, padding=16)
    marco.grid(sticky="nsew")

    ttk.Label(marco, text="Generador de Actas de Plan Concertado",
              font=("Segoe UI", 13, "bold")).grid(row=0, column=0, columnspan=2,
                                                  sticky="w", pady=(0, 12))

    v = {}
    campos = [
        ("colegio", "Colegio"), ("grado", "Grado"), ("trimestre", "Trimestre"),
        ("ficha", "N° de ficha"), ("vocero", "Vocero de la ficha"),
        ("instructor", "Instructor"), ("acta_no", "Acta N°"),
        ("fecha", "Fecha de la reunión (dd/mm/aaaa)"),
        ("hora_inicio", "Hora inicio"), ("hora_fin", "Hora fin"),
        ("dia", "Día(s) de sesión"), ("ambiente", "Ambiente de formación"),
    ]
    for clave, _ in campos:
        v[clave] = tk.StringVar()

    fila = 1
    widgets = {}
    for clave, etiqueta in campos:
        ttk.Label(marco, text=etiqueta).grid(row=fila, column=0, sticky="w", pady=3)
        if clave == "colegio":
            w = ttk.Combobox(marco, textvariable=v[clave], width=42, state="readonly",
                             values=[c["nombre"] for c in cfg["colegios"]])
        elif clave == "grado":
            w = ttk.Combobox(marco, textvariable=v[clave], width=42, state="readonly")
        elif clave == "trimestre":
            w = ttk.Combobox(marco, textvariable=v[clave], width=42, state="readonly",
                             values=["1", "2", "3", "4"])
        else:
            w = ttk.Entry(marco, textvariable=v[clave], width=44)
        w.grid(row=fila, column=1, sticky="w", pady=3)
        widgets[clave] = w
        fila += 1

    v["instructor"].set(cfg.get("instructor", ""))
    v["fecha"].set(date.today().strftime("%d/%m/%Y"))
    v["trimestre"].set("1")

    def colegio_actual():
        return next((c for c in cfg["colegios"] if c["nombre"] == v["colegio"].get()), None)

    def al_cambiar_colegio(*_):
        col = colegio_actual()
        if not col:
            return
        grados = [f["grado"] for f in col["fichas"]]
        widgets["grado"]["values"] = grados
        if v["grado"].get() not in grados and grados:
            v["grado"].set(grados[0])
        v["ambiente"].set(col.get("ambiente", ""))
        al_cambiar_grado()

    def al_cambiar_grado(*_):
        col = colegio_actual()
        if not col:
            return
        f = next((x for x in col["fichas"] if str(x["grado"]) == v["grado"].get()), None)
        if not f:
            return
        v["ficha"].set(f.get("ficha", ""))
        v["vocero"].set(f.get("vocero", ""))
        v["dia"].set(f.get("dia", ""))
        v["hora_inicio"].set(f.get("hora_inicio", ""))
        v["hora_fin"].set(f.get("hora_fin", ""))

    widgets["colegio"].bind("<<ComboboxSelected>>", al_cambiar_colegio)
    widgets["grado"].bind("<<ComboboxSelected>>", al_cambiar_grado)
    if cfg["colegios"]:
        v["colegio"].set(cfg["colegios"][0]["nombre"])
        al_cambiar_colegio()

    estado = ttk.Label(marco, text="", foreground="#1a7f37", wraplength=430,
                       justify="left")
    estado.grid(row=fila + 1, column=0, columnspan=2, sticky="w", pady=(10, 0))

    def leer_fecha():
        try:
            return datetime.strptime(v["fecha"].get().strip(), "%d/%m/%Y").date()
        except ValueError:
            return None

    def datos_actuales(col, ficha_cfg):
        return {
            "colegio": col["nombre"], "ambiente": v["ambiente"].get(),
            "grado": ficha_cfg["grado"], "ficha": ficha_cfg.get("ficha", ""),
            "trimestre": v["trimestre"].get(), "vocero": ficha_cfg.get("vocero", ""),
            "instructor": v["instructor"].get(), "acta_no": v["acta_no"].get(),
            "fecha": leer_fecha(), "dia": ficha_cfg.get("dia", ""),
            "hora_inicio": ficha_cfg.get("hora_inicio", ""),
            "hora_fin": ficha_cfg.get("hora_fin", ""),
        }

    def generar_uno():
        col = colegio_actual()
        if not col:
            messagebox.showwarning("Falta información", "Seleccione un colegio.")
            return
        if leer_fecha() is None:
            messagebox.showwarning("Fecha inválida", "Use el formato dd/mm/aaaa.")
            return
        if not v["ficha"].get().strip():
            if not messagebox.askyesno(
                    "Ficha vacía",
                    "No escribió el número de ficha. ¿Genera el acta igual, "
                    "dejando puntos suspensivos?"):
                return
        datos = {
            "colegio": col["nombre"], "ambiente": v["ambiente"].get(),
            "grado": v["grado"].get(), "ficha": v["ficha"].get(),
            "trimestre": v["trimestre"].get(), "vocero": v["vocero"].get(),
            "instructor": v["instructor"].get(), "acta_no": v["acta_no"].get(),
            "fecha": leer_fecha(), "dia": v["dia"].get(),
            "hora_inicio": v["hora_inicio"].get(), "hora_fin": v["hora_fin"].get(),
        }
        try:
            ruta = generar_acta(cfg, datos)
        except Exception as exc:
            messagebox.showerror("Error al generar el acta", str(exc))
            return
        estado.config(text="Acta generada:\n%s" % ruta)
        abrir_carpeta(os.path.dirname(ruta))

    def generar_todas():
        if leer_fecha() is None:
            messagebox.showwarning("Fecha inválida", "Use el formato dd/mm/aaaa.")
            return
        if not messagebox.askyesno(
                "Generar todas",
                "Se generará un acta por cada ficha registrada en el archivo de "
                "configuración, para el trimestre %s.\n¿Continuar?" % v["trimestre"].get()):
            return
        hechas, errores = [], []
        for col in cfg["colegios"]:
            for f in col["fichas"]:
                datos = datos_actuales(col, f)
                datos["ambiente"] = col.get("ambiente", "")
                try:
                    hechas.append(generar_acta(cfg, datos))
                except Exception as exc:
                    errores.append("%s grado %s: %s" % (col["nombre"], f["grado"], exc))
        texto = "Se generaron %d actas." % len(hechas)
        if errores:
            texto += "\nCon errores:\n" + "\n".join(errores)
        estado.config(text=texto)
        if hechas:
            abrir_carpeta(os.path.dirname(hechas[0]))

    botones = ttk.Frame(marco)
    botones.grid(row=fila, column=0, columnspan=2, sticky="w", pady=(14, 0))
    ttk.Button(botones, text="Generar acta", command=generar_uno).grid(row=0, column=0)
    ttk.Button(botones, text="Generar todas las fichas",
               command=generar_todas).grid(row=0, column=1, padx=8)
    ttk.Button(botones, text="Salir", command=ventana.destroy).grid(row=0, column=2)

    ventana.mainloop()


def abrir_carpeta(ruta):
    try:
        if sys.platform.startswith("win"):
            os.startfile(ruta)  # noqa: S606
        elif sys.platform == "darwin":
            os.system('open "%s"' % ruta)
        else:
            os.system('xdg-open "%s" >/dev/null 2>&1 &' % ruta)
    except Exception:
        pass


if __name__ == "__main__":
    main()
